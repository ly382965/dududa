"""USTC homework APIs used by the university's wdzyappustc frontend."""

from __future__ import annotations

import ast
import io
import json
import mimetypes
import re
import zipfile
from html.parser import HTMLParser
from pathlib import PurePosixPath
from urllib.parse import urljoin, urlsplit, urlunsplit

import httpx

ORIGIN = "https://yjs1.ustc.edu.cn"
APP_PATH = "/gsapp/sys/wdzyappustc"
APP_URL = ORIGIN + APP_PATH + "/*default/index.do"
FILE_PATH = "/gsapp/sys/emapcomponent/file"
MAX_FILE_BYTES = 20 * 1024 * 1024


class HomeworkError(RuntimeError):
    pass


class TextParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []

    def handle_data(self, data):
        self.parts.append(data)


def text_from_file(name: str, content: bytes) -> str:
    extension = PurePosixPath(name).suffix.lower()
    if extension in {".txt", ".md", ".csv", ".py", ".tex"}:
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
    elif extension == ".pdf":
        from pypdf import PdfReader

        return "\n".join(
            page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages
        )
    elif extension == ".docx":
        from docx import Document

        document = Document(io.BytesIO(content))
        return "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [
                " | ".join(cell.text for cell in row.cells)
                for table in document.tables
                for row in table.rows
            ]
        )
    elif extension in {".html", ".htm"}:
        parser = TextParser()
        parser.feed(content.decode("utf-8"))
        return "\n".join(parser.parts)
    elif extension == ".zip":
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_FILE_BYTES:
                raise HomeworkError("attachment_too_large")
            texts = []
            for item in archive.infolist():
                if item.is_dir():
                    continue
                if item.filename.lower().endswith(".zip"):
                    return ""
                text = text_from_file(item.filename, archive.read(item))
                if not text.strip():
                    return ""
                texts.append(f"### {PurePosixPath(item.filename).name}\n{text}")
            return "\n\n".join(texts)
    return ""


def assignment_requirements(item: dict) -> dict:
    return {
        key: item.get(key)
        for key in (
            "id",
            "title",
            "course",
            "description",
            "deadline",
            "requirements_token",
        )
    }


class HomeworkClient:
    def __init__(self, session: httpx.AsyncClient):
        self.session = session
        self.student_id = ""
        self.upload_types = []
        self.headers = {
            "X-Requested-With": "XMLHttpRequest",
            "Referer": APP_URL,
            "Origin": ORIGIN,
        }

    async def initialize(self):
        url = APP_URL
        for _ in range(8):
            if "/yjsrzfwapp/dbLogin/" in urlsplit(url).path:
                raise HomeworkError("homework_session_expired")
            response = await self.session.get(url, follow_redirects=False)
            if not response.is_redirect:
                break
            parts = urlsplit(urljoin(url, response.headers.get("location", "")))
            if parts.hostname != "yjs1.ustc.edu.cn" or not parts.path.startswith(
                "/gsapp/"
            ):
                raise HomeworkError("homework_session_expired")
            url = urlunsplit(("https", parts.hostname, parts.path, parts.query, ""))
        else:
            raise HomeworkError("homework_initialization_failed")
        if response.status_code != 200:
            raise HomeworkError("homework_initialization_failed")
        match = re.search(r"\bvar\s+pageMeta\s*=\s*(\{[^\n]+\});", response.text)
        if not match:
            raise HomeworkError("homework_identity_missing")
        student_id = json.loads(match.group(1)).get("params", {}).get("USERID")
        if (
            not isinstance(student_id, str)
            or not student_id.strip()
            or student_id.strip().lower() in {"null", "none", "undefined"}
        ):
            raise HomeworkError("homework_identity_missing")
        self.student_id = student_id.strip()
        schema = await self.post(APP_PATH + "/modules/wdzy.do", {"*json": "1"})
        model = next(
            model for model in schema["models"] if model["name"] == "wdzycx_form"
        )
        field = next(field for field in model["controls"] if field["name"] == "ZY")
        # The school serves this JSONParam as a Python-compatible single-quoted literal.
        options = ast.literal_eval(field.get("form.JSONParam", "{}"))
        self.upload_types = options.get("type", [])

    async def post(self, path: str, data: dict) -> dict:
        response = await self.session.post(
            path, data=data, headers=self.headers, follow_redirects=False
        )
        if not response.is_success:
            raise HomeworkError("homework_request_failed")
        try:
            payload = response.json()
        except ValueError:
            raise HomeworkError("homework_response_invalid") from None
        if (
            not isinstance(payload, dict)
            or payload.get("success") is False
            or payload.get("code") not in (None, 0, "0")
        ):
            raise HomeworkError("homework_request_rejected")
        return payload

    async def table(self, action: str, data: dict) -> dict:
        payload = await self.post(
            APP_PATH + f"/modules/wdzy/{action}.do", {"XH": self.student_id, **data}
        )
        table = payload.get("datas", {}).get(action)
        if not isinstance(table, dict) or not isinstance(table.get("rows"), list):
            raise HomeworkError("homework_table_invalid")
        if table.get("extParams", {}).get("code", 1) not in (1, "1"):
            raise HomeworkError("homework_query_rejected")
        return table

    async def list_rows(self) -> list[dict]:
        rows = {}
        total = None
        for page in range(1, 101):
            table = await self.table(
                "wdzycx", {"pageNumber": page, "pageSize": 100, "querySetting": "[]"}
            )
            current_total = int(table["totalSize"])
            if total is not None and total != current_total:
                raise HomeworkError("homework_list_changed_during_fetch")
            total = current_total
            for row in table["rows"]:
                if row["WID"] in rows:
                    raise HomeworkError("homework_repeated_page")
                rows[row["WID"]] = row
            if len(rows) == total:
                return list(rows.values())
            if not table["rows"] or len(rows) > total:
                break
        raise HomeworkError("homework_list_incomplete")

    async def snapshot(self):
        rows = await self.list_rows()
        pending, submitted = [], []
        for row in rows:
            # The list omits submission fields; verify each student's actual detail.
            item = await self.detail(row["WID"], attachments=False)
            if item["submitted"]:
                submitted.append(item)
            else:
                if item["requirements_token"]:
                    item["attachments"] = await self.read_attachments(
                        item["requirements_token"]
                    )
                pending.append(item)
        return {
            "total": len(rows),
            "assignments": pending,
            "submitted_assignments": submitted,
        }

    async def detail(self, assignment_id: str, *, attachments: bool = True) -> dict:
        table = await self.table(
            "wdzycx_form", {"WID": assignment_id, "pageNumber": 1, "pageSize": 1}
        )
        if len(table["rows"]) != 1 or table["rows"][0]["WID"] != assignment_id:
            raise HomeworkError("homework_not_found")
        row = table["rows"][0]
        if str(row.get("XH")) != self.student_id:
            raise HomeworkError("homework_owner_mismatch")
        edit = await self.post(
            APP_PATH + "/wdzy/validateCanEdit.do",
            {"WID": assignment_id, "XH": self.student_id},
        )
        item = {
            "id": assignment_id,
            "title": row.get("ZYMC") or assignment_id,
            "course": row.get("KCMC", ""),
            "class_name": row.get("BJMC", ""),
            "description": row.get("ZYSM") or "",
            "starts_at": row.get("ZYKSSJ"),
            "deadline": row.get("ZYJZSJ"),
            "requirements_token": row.get("ZYYQFJ") or "",
            "submission_token": row.get("ZY") or "",
            "submission_id": row.get("XSZYWID") or "",
            "submitted": bool(row.get("ZY")),
            "can_edit": edit.get("canEdit") is True,
            "edit_reason": edit.get("msg") or "",
            "upload_types": self.upload_types,
            "attachments": [],
        }
        if attachments and item["requirements_token"]:
            item["attachments"] = await self.read_attachments(
                item["requirements_token"]
            )
        return item

    async def read_attachments(self, token: str) -> list[dict]:
        payload = await self.post(
            FILE_PATH + "/getUploadedAttachment.do", {"fileToken": token}
        )
        files = payload.get("items")
        if not isinstance(files, list):
            raise HomeworkError("homework_attachment_list_invalid")
        results = []
        for item in files:
            name = item.get("name") or item.get("fileName") or "attachment"
            parts = urlsplit(urljoin(ORIGIN, item.get("fileUrl") or ""))
            if parts.hostname != "yjs1.ustc.edu.cn" or not parts.path.startswith(
                FILE_PATH + "/"
            ):
                raise HomeworkError("homework_attachment_url_invalid")
            url = urlunsplit(("https", parts.hostname, parts.path, parts.query, ""))
            async with self.session.stream(
                "GET", url, headers=self.headers, follow_redirects=False
            ) as response:
                if not response.is_success:
                    raise HomeworkError("homework_attachment_download_failed")
                content = bytearray()
                async for chunk in response.aiter_bytes():
                    content.extend(chunk)
                    if len(content) > MAX_FILE_BYTES:
                        raise HomeworkError("attachment_too_large")
            try:
                text = text_from_file(name, bytes(content))
            except Exception:  # noqa: BLE001 - an unreadable attachment must not hide its assignment
                text = ""
            results.append(
                {
                    "name": name,
                    "text": text[:60000],
                    "readable": bool(text.strip()) and len(text) <= 60000,
                }
            )
        return results

    async def submit(
        self, expected: dict, filename: str, content: bytes, token: str
    ) -> dict:
        current = await self.detail(expected["id"], attachments=False)
        if current["submission_token"] == token:
            return {"status": "submitted", "submission_id": current["submission_id"]}
        if current["submitted"]:
            raise HomeworkError("homework_already_submitted")
        if assignment_requirements(current) != assignment_requirements(expected):
            raise HomeworkError("homework_requirements_changed")
        if not current["can_edit"]:
            raise HomeworkError("homework_submission_closed")
        extension = PurePosixPath(filename).suffix.lstrip(".").lower()
        if self.upload_types and extension not in self.upload_types:
            raise HomeworkError("homework_file_type_rejected")
        scope = token[:-1]
        response = await self.session.post(
            FILE_PATH + "/uploadTempFile.do",
            headers=self.headers,
            data={
                "scope": scope,
                "fileToken": token,
                "storeId": "file",
                "isSingle": "0",
            },
            files={
                "files[]": (
                    filename,
                    content,
                    mimetypes.guess_type(filename)[0] or "application/octet-stream",
                )
            },
            follow_redirects=False,
        )
        if not response.is_success or response.json().get("success") is not True:
            raise HomeworkError("homework_upload_failed")
        saved = await self.post(
            FILE_PATH + f"/saveAttachment/{scope}/{token}.do",
            {"fileToken": token, "scope": scope},
        )
        if saved.get("success") is not True:
            raise HomeworkError("homework_attachment_save_failed")
        payload = await self.post(
            APP_PATH + "/modules/wdzy/sczy.do",
            {
                "WID": current["submission_id"],
                "ZYWID": current["id"],
                "XH": self.student_id,
                "ZY": token,
                "CJ": "99",
            },
        )
        if payload.get("datas", {}).get("sczy", {}).get("code") not in (1, "1"):
            raise HomeworkError("homework_save_rejected")
        try:
            receipt = await self.detail(current["id"], attachments=False)
        except HomeworkError:
            raise HomeworkError("homework_submission_unconfirmed") from None
        if receipt["submission_token"] != token:
            raise HomeworkError("homework_submission_unconfirmed")
        return {"status": "submitted", "submission_id": receipt["submission_id"]}
