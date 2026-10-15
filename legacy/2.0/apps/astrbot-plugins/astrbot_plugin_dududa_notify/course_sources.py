"""Read-only course pages and files. Never opens a test or starts an attempt."""

from __future__ import annotations

import asyncio
import hashlib
import io
import json
import re
import shutil
import subprocess
import tempfile
from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone
from email.utils import parsedate_to_datetime
from pathlib import Path
from urllib.parse import parse_qs, urlencode, urljoin, urlsplit, urlunsplit

import httpx
from lxml import html

if __package__:
    from .course_catalog import (
        ALGORITHM_URL,
        BB_ORIGIN,
        DEFAULT_COURSES,
        NETWORK_URL,
        validate_courses,
    )
else:
    from course_catalog import (
        ALGORITHM_URL,
        BB_ORIGIN,
        DEFAULT_COURSES,
        NETWORK_URL,
        validate_courses,
    )

CHINA_TIME = timezone(timedelta(hours=8))
MAX_FILE_BYTES = 20 * 1024 * 1024


class CourseError(ValueError):
    pass


def document(content):
    doc = html.fromstring(
        content.decode("utf-8", errors="replace")
        if isinstance(content, bytes)
        else content
    )
    for node in doc.xpath(
        "//script | //style | //*[contains(@class,'contextMenuContainer')]"
    ):
        node.drop_tree()
    return doc


def plain(node):
    return " ".join(node.text_content().split())


def explicit_deadline(text, *, default_year=None):
    # Dates must be tied to a deadline label; a lecture date is not a due date.
    pattern = r"(?:截止(?:时间|日期)?|提交时间|提交日期|due(?:\s+date)?|ddl)\s*[:：为是]?\s*(?:(20\d{2})\s*[年/.-]\s*)?(\d{1,2})\s*[月/.-]\s*(\d{1,2})\s*日?(?:\s*(?:星期[一二三四五六日天]|\([^)]*\)))?\s*(?:(\d{1,2})[:：](\d{2}))?"
    matches = list(re.finditer(pattern, text, re.IGNORECASE))
    values = set()
    for match in matches:
        year, month, day, hour, minute = match.groups()
        year = year or default_year
        if year is None:
            continue
        try:
            value = datetime(
                int(year),
                int(month),
                int(day),
                int(hour or 0),
                int(minute or 0),
                tzinfo=CHINA_TIME,
            )
        except ValueError:
            continue
        values.add(
            (
                value.isoformat() if hour else value.date().isoformat(),
                "minute" if hour else "date",
            )
        )
    if len(values) != 1:
        return None, "unknown"
    return values.pop()


def task_item(key, title, text, url, *, default_year=None):
    deadline, precision = explicit_deadline(text, default_year=default_year)
    return {
        "id": key,
        "title": title,
        "description": text,
        "url": url,
        "deadline": deadline,
        "precision": precision,
        "submitted": None,
    }


async def bb_get(client, url, *, method="GET", headers=None):
    for _ in range(8):
        parts = urlsplit(url)
        if (
            parts.hostname != "www.bb.ustc.edu.cn"
            or parts.username
            or parts.port not in (None, 443)
        ):
            raise CourseError("course_redirect_rejected")
        url = urlunsplit(("https", parts.netloc, parts.path, parts.query, ""))
        response = await client.request(method, url, headers=headers)
        if not response.is_redirect:
            return response
        url = urljoin(url, response.headers.get("location", ""))
    raise CourseError("course_redirect_limit")


@asynccontextmanager
async def blackboard_session(credentials, timeout=45):
    from pyustc.cas import CASClient

    if __package__:
        from .campus_auth import graduate_ticket
    else:
        from campus_auth import graduate_ticket

    async with httpx.AsyncClient(
        trust_env=False, timeout=timeout, follow_redirects=False
    ) as client:
        async with asyncio.timeout(timeout * 2):
            response = await bb_get(client, BB_ORIGIN + "/webapps/login/")
            response.raise_for_status()
            doc = document(response.content)
            links = doc.xpath("//a[contains(.,'统一身份认证')]/@href")
            if not links:
                raise CourseError("blackboard_login_link_missing")
            login = urlsplit(links[0])
            if (
                login.scheme != "https"
                or login.hostname != "passport.ustc.edu.cn"
                or login.path != "/login"
            ):
                raise CourseError("blackboard_login_link_invalid")
            response = await client.get(links[0])
            redirect = urlsplit(response.headers.get("location", ""))
            if redirect.hostname != "id.ustc.edu.cn" or redirect.path != "/cas/login":
                raise CourseError("blackboard_cas_redirect_invalid")
            service = parse_qs(redirect.query).get("service", [""])[0]
            target = urlsplit(service)
            if (
                target.hostname != "www.bb.ustc.edu.cn"
                or target.scheme != "https"
                or target.path
                != "/webapps/bb-SSOIntegrationDemo-BBLEARN/execute/authValidate/customLogin"
            ):
                raise CourseError("blackboard_ticket_target_invalid")
            cas = CASClient.login_by_pwd(credentials.username, credentials.password)
            cas._client.timeout = httpx.Timeout(timeout)
            async with cas:
                ticket = await graduate_ticket(cas, service)
            response = await bb_get(
                client, service + "&" + urlencode({"ticket": ticket})
            )
            response.raise_for_status()
        yield client


def parse_blackboard(doc, page_url, course_id):
    if not doc.xpath('//*[@id="courseMenuPalette_contents"]'):
        raise CourseError("blackboard_session_or_access_failed")
    tasks, materials, folders, notices = [], [], [], []
    for node in doc.xpath('//*[@id="content_listContainer"]/li'):
        titles = node.xpath(".//h3")
        title = plain(titles[0]) if titles else "未命名资源"
        key = node.get("id", "")
        text = plain(node)
        is_folder = bool(node.xpath('.//h3//a[contains(@href,"/listContent.jsp")]'))
        if not is_folder and re.search(
            r"作业|实验|小测|测验|homework|assignment|\blab\d*\b|project|quiz",
            title,
            re.IGNORECASE,
        ):
            tasks.append(task_item(key, title, text, page_url))
        for anchor in node.xpath(".//a[@href]"):
            url = urljoin(page_url, anchor.get("href"))
            parts = urlsplit(url)
            if parts.hostname != "www.bb.ustc.edu.cn":
                continue
            query = parse_qs(parts.query)
            if parts.path.endswith("/listContent.jsp") and query.get("course_id") == [
                course_id
            ]:
                folders.append(url)
            elif parts.path.startswith("/bbcswebdav/"):
                materials.append(
                    {
                        "id": parts.path,
                        "title": plain(anchor) or title,
                        "url": url,
                        "parent_id": key,
                        "parent_title": title,
                        "description": text,
                    }
                )
    for anchor in doc.xpath('//*[@id="courseMenuPalette_contents"]//a[@href]'):
        url = urljoin(page_url, anchor.get("href"))
        parts = urlsplit(url)
        if parts.hostname != "www.bb.ustc.edu.cn" or parse_qs(parts.query).get(
            "course_id"
        ) != [course_id]:
            continue
        if parts.path.endswith("/listContent.jsp"):
            folders.append(url)
        elif "课程通知" in plain(anchor) and parts.path.endswith("/launchLink.jsp"):
            notices.append(url)
    return tasks, materials, folders, notices


def parse_algorithms(doc, page_url):
    if not doc.xpath('//*[@id="homework"]') or not doc.xpath('//*[@id="courseware"]'):
        raise CourseError("course_website_layout_changed")
    tasks, materials = [], []
    sections = doc.xpath('//*[@id="homework"]')
    for section in sections:
        text = plain(section)
        for index, match in enumerate(
            re.finditer(r"第[一二三四五六七八九十\d]+次作业[^。\n]*", text), 1
        ):
            tasks.append(
                task_item(
                    f"homework-{index}", match.group(0), text, page_url + "#homework"
                )
            )
    for anchor in doc.xpath('//*[@id="courseware"]//a[@href]'):
        url = urljoin(page_url, anchor.get("href"))
        if url.startswith(ALGORITHM_URL + "ppt/"):
            materials.append(
                {
                    "id": urlsplit(url).path,
                    "title": plain(anchor),
                    "url": url,
                    "parent_title": "课程课件",
                    "description": "",
                }
            )
    return tasks, materials


def parse_networks(doc, page_url):
    notices = doc.xpath(
        '//div[contains(concat(" ",normalize-space(@class)," ")," paper-report-notice ")]'
    )
    schedules = doc.xpath(
        '//div[contains(concat(" ",normalize-space(@class)," ")," course-table--schedule ")]'
    )
    headings = doc.xpath("//article/h1")
    if not notices or not schedules or not headings:
        raise CourseError("course_website_layout_changed")
    year = re.search(r"\b(20\d{2})\b", plain(headings[0]))
    if year is None:
        raise CourseError("network_course_year_missing")
    year = int(year.group(1))
    tasks, materials = [], []
    titles = {
        "线上汇报": ("paper-online", "论文汇报：线上方案（二选一）"),
        "线下汇报": ("paper-offline", "论文汇报：线下方案（二选一，课前提交）"),
        "第一部分": ("lab-tasks", "选做实验：前四次任务（代码、实验报告）"),
        "第二部分": ("lab-leaderboard", "选做实验：平台打榜（须先交前四次任务）"),
    }
    for notice in notices:
        for node in notice.xpath("./ul/li"):
            labels = node.xpath("./p[1]")
            if not labels:
                continue
            label, text = plain(labels[0]), plain(node)
            spec = next((v for k, v in titles.items() if label.startswith(k)), None)
            if spec is None:
                headings = labels[0].xpath("./strong")
                if (
                    not headings
                    or not label.startswith(plain(headings[0]))
                    or not re.search(r"作业|任务|实验|项目|测验", label)
                ):
                    continue
                spec = ("task:" + label, label)
            item = task_item(*spec, text, page_url + "#_1", default_year=year)
            if spec[0].startswith("lab-"):
                ends = re.findall(r"实验时间[：:]\s*[^；;]*?[至到]\s*([^；;]+)", text)
                item["deadline"], item["precision"] = explicit_deadline(
                    " ".join("截止日期：" + end for end in ends), default_year=year
                )
            tasks.append(item)
    for schedule in schedules:
        for row in schedule.xpath(".//tbody/tr"):
            cells = row.xpath("./td")
            if len(cells) < 4:
                continue
            description = plain(cells[0]) + " " + plain(cells[1])
            for anchor in cells[3].xpath(".//a[@href]"):
                url = urljoin(page_url, anchor.get("href"))
                if not url.startswith(NETWORK_URL + "slides/"):
                    continue
                materials.append(
                    {
                        "id": urlsplit(url).path,
                        "title": description + "：" + plain(anchor),
                        "url": url,
                        "parent_title": "课程课件",
                        "description": description,
                    }
                )
    return tasks, materials


def file_text(name, content):
    if __package__:
        from .assignment_client import text_from_file
    else:
        from assignment_client import text_from_file
    extension = Path(name).suffix.lower()
    if extension == ".pptx":
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        presentation = Presentation(io.BytesIO(content))

        def shape_text(shapes):
            for shape in shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    yield from shape_text(shape.shapes)
                elif shape.has_text_frame:
                    yield shape.text
                elif shape.has_table:
                    for row in shape.table.rows:
                        yield " | ".join(cell.text for cell in row.cells)

        return "\n\n".join(
            f"第 {i} 页\n" + "\n".join(shape_text(slide.shapes))
            for i, slide in enumerate(presentation.slides, 1)
        )
    if extension in {".ppt", ".doc"}:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / ("material" + extension)
            path.write_bytes(content)
            if extension == ".ppt" and shutil.which("libreoffice"):
                subprocess.run(
                    [
                        "libreoffice",
                        "-env:UserInstallation="
                        + (Path(directory) / "profile").as_uri(),
                        "--headless",
                        "--convert-to",
                        "pptx",
                        "--outdir",
                        directory,
                        str(path),
                    ],
                    capture_output=True,
                    timeout=60,
                    check=True,
                )
                converted = path.with_suffix(".pptx")
                if not converted.is_file():
                    raise CourseError("legacy_ppt_conversion_failed")
                return file_text(converted.name, converted.read_bytes())
            result = subprocess.run(
                [
                    "catppt" if extension == ".ppt" else "catdoc",
                    "-d",
                    "utf-8",
                    str(path),
                ],
                capture_output=True,
                timeout=30,
                check=True,
            )
            text = result.stdout.decode("utf-8", errors="replace")
            if extension == ".ppt" and not re.sub(r"\d+_默认设计模板|\s+", "", text):
                raise CourseError("legacy_ppt_converter_required")
            return text
    return text_from_file(name, content)


def modified_date(value):
    try:
        return (
            parsedate_to_datetime(value)
            .astimezone(CHINA_TIME)
            .isoformat(timespec="seconds")
        )
    except (TypeError, ValueError):
        return None


def parse_announcements(doc, url):
    if not doc.xpath('//*[@id="announcementList"]'):
        raise CourseError("blackboard_announcements_unreadable")
    result = []
    for node in doc.xpath('//*[@id="announcementList"]/li'):
        text = plain(node)
        if not re.search(
            r"作业|实验|小测|测验|截止|homework|assignment|\blab\b|due|deadline",
            text,
            re.IGNORECASE,
        ):
            continue
        titles = node.xpath(".//h3 | .//h2")
        title = plain(titles[0]) if titles else text[:70]
        result.append(
            task_item("notice:" + (node.get("id") or title), title, text, url)
        )
    return result


async def material_snapshot(client, material, previous, *, blackboard):
    url = material["url"]
    network_file = not blackboard and url.startswith(NETWORK_URL + "slides/")

    async def request(method):
        if blackboard:
            return await bb_get(client, url, method=method)
        if not (url.startswith(ALGORITHM_URL + "ppt/") or network_file):
            raise CourseError("course_file_url_rejected")
        return await client.request(method, url)

    head = await request("HEAD")
    if head.status_code not in (200, 405, 501):
        head.raise_for_status()
    metadata = {
        key: head.headers.get(key, "")
        for key in ("etag", "last-modified", "content-length", "content-type")
    }
    signature = json.dumps(
        {
            "metadata": metadata,
            "title": material["title"],
            "description": material["description"],
        },
        sort_keys=True,
        ensure_ascii=False,
    )
    snapshot = {
        **{
            key: value
            for key, value in material.items()
            if key in {"id", "title", "url", "parent_id", "parent_title", "description"}
        },
        "http_signature" if network_file else "signature": signature,
        "modified_at": None
        if network_file
        else modified_date(metadata["last-modified"]),
    }
    if (
        previous
        and previous.get("http_signature" if network_file else "signature") == signature
        and not previous.get("read_error")
        and (metadata["etag"] or metadata["last-modified"])
    ):
        return {**previous, **snapshot}
    if (
        metadata["content-length"].isdigit()
        and int(metadata["content-length"]) > MAX_FILE_BYTES
    ):
        return {**snapshot, "text": "", "read_error": "文件超过20MB，需人工阅读"}
    response = await request("GET")
    response.raise_for_status()
    if network_file:
        # GitHub Pages ETags include build time; only changed file bytes are updates.
        snapshot["signature"] = hashlib.sha256(response.content).hexdigest()
    if len(response.content) > MAX_FILE_BYTES:
        return {**snapshot, "text": "", "read_error": "文件超过20MB，需人工阅读"}
    extension = (
        Path(urlsplit(url).path).suffix.lower()
        or Path(material["title"]).suffix.lower()
    )
    content_type = response.headers.get("content-type", "").split(";")[0]
    extension = {
        "application/pdf": ".pdf",
        "application/vnd.ms-powerpoint": ".ppt",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
        "application/msword": ".doc",
    }.get(content_type, extension)
    try:
        text = file_text("courseware" + extension, response.content)
    except Exception as exc:  # noqa: BLE001 - unsupported or scanned files must stay visible
        return {
            **snapshot,
            "text": "",
            "read_error": "课件解析失败：" + type(exc).__name__,
        }
    if not text.strip():
        return {
            **snapshot,
            "text": "",
            "read_error": "无法提取正文（可能是扫描件），需人工阅读",
        }
    normalized = text.encode("utf-8", errors="replace").decode("utf-8")
    return {
        **snapshot,
        "text": normalized,
        "read_error": "",
        "read_warning": "部分符号无法解码，请对照原课件" if normalized != text else "",
    }


async def read_blackboard(client, course, cache):
    cid = parse_qs(urlsplit(course["url"]).query)["course_id"][0]
    todo, seen, tasks, materials, notice_urls = [course["url"]], set(), {}, {}, set()
    while todo:
        url = todo.pop(0)
        content_id = parse_qs(urlsplit(url).query).get("content_id", [""])[0]
        if content_id in seen:
            continue
        if len(seen) >= 100:
            raise CourseError("course_folder_limit")
        seen.add(content_id)
        response = await bb_get(client, url)
        response.raise_for_status()
        found_tasks, files, folders, notices = parse_blackboard(
            document(response.content), url, cid
        )
        tasks.update((item["id"], item) for item in found_tasks)
        materials.update((item["id"], item) for item in files)
        todo.extend(folders)
        notice_urls.update(notices)
    for url in notice_urls:
        response = await bb_get(client, url)
        response.raise_for_status()
        for task in parse_announcements(document(response.content), str(response.url)):
            tasks[task["id"]] = task
    for key, material in list(materials.items()):
        materials[key] = await material_snapshot(
            client, material, cache.get(key), blackboard=True
        )
        parent = material.get("parent_id")
        if parent in tasks or re.search(
            r"作业|实验|homework|assignment|\blab\d*\b",
            material["title"],
            re.IGNORECASE,
        ):
            task_key = parent if parent in tasks else key
            title = tasks[task_key]["title"] if task_key in tasks else material["title"]
            text = material["description"] + "\n" + materials[key].get("text", "")
            tasks[task_key] = task_item(task_key, title, text, material["url"])
    return {
        "id": course["id"],
        "status": "ok",
        "tasks": list(tasks.values()),
        "materials": list(materials.values()),
    }


async def fetch_courses(request, credentials):
    courses = validate_courses(request.get("courses") or DEFAULT_COURSES)
    caches = request.get("cache", {})
    results = []
    for course in courses:
        if course["kind"] in {"pending", "graduate"}:
            continue
        if course["kind"] == "website":
            try:
                async with httpx.AsyncClient(
                    trust_env=False,
                    timeout=request.get("timeout", 45),
                    follow_redirects=False,
                ) as client:
                    response = await client.get(course["url"])
                    response.raise_for_status()
                    parser = (
                        parse_networks
                        if course["url"] == NETWORK_URL
                        else parse_algorithms
                    )
                    tasks, materials = parser(document(response.content), course["url"])
                    files = [
                        await material_snapshot(
                            client,
                            material,
                            caches.get(course["id"], {}).get(material["id"]),
                            blackboard=False,
                        )
                        for material in materials
                    ]
                    results.append(
                        {
                            "id": course["id"],
                            "status": "ok",
                            "tasks": tasks,
                            "materials": files,
                        }
                    )
            except Exception as exc:  # noqa: BLE001 - isolate source outages and sanitize all network errors
                results.append(
                    {
                        "id": course["id"],
                        "status": "error",
                        "error": str(exc)
                        if isinstance(exc, CourseError)
                        else type(exc).__name__,
                    }
                )
    bb_courses = [course for course in courses if course["kind"] == "blackboard"]
    if bb_courses:
        try:
            async with blackboard_session(
                credentials, request.get("timeout", 45)
            ) as client:
                for course in bb_courses:
                    try:
                        results.append(
                            await read_blackboard(
                                client, course, caches.get(course["id"], {})
                            )
                        )
                    except Exception as exc:  # noqa: BLE001 - do not replace an inaccessible course with an empty success
                        results.append(
                            {
                                "id": course["id"],
                                "status": "error",
                                "error": str(exc)
                                if isinstance(exc, CourseError)
                                else type(exc).__name__,
                            }
                        )
        except Exception as exc:  # noqa: BLE001 - CAS errors may contain secrets
            results.extend(
                {"id": course["id"], "status": "error", "error": type(exc).__name__}
                for course in bb_courses
                if not any(r["id"] == course["id"] for r in results)
            )
    return {"courses": results}
