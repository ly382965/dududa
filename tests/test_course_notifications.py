from __future__ import annotations

import io
import json
import tempfile
import unittest
from datetime import datetime, timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock, patch

import httpx
from astrbot_plugin_dududa_notify.course_catalog import (
    DEFAULT_COURSES,
    NETWORK_URL,
    validate_courses,
)
from astrbot_plugin_dududa_notify.course_report import render_report
from astrbot_plugin_dududa_notify.course_sources import (
    BB_ORIGIN,
    CourseError,
    bb_get,
    document,
    explicit_deadline,
    fetch_courses,
    file_text,
    material_snapshot,
    parse_algorithms,
    parse_announcements,
    parse_blackboard,
    parse_networks,
)
from astrbot_plugin_dududa_notify.course_tasks import (
    CourseService,
    deadline_label,
    report_text,
)
from astrbot_plugin_dududa_notify.homework import CHINA_TIME

from tests.test_private_notifications import RECIPIENT, load_plugin

COURSE = {
    "id": "ml",
    "name": "机器学习",
    "kind": "blackboard",
    "url": DEFAULT_COURSES[2]["url"],
}

NETWORK_PAGE = """<article><h1>USTC 高级计算机网络（2026）课程主页</h1>
<div class="paper-report-notice">
<ul>
<li><p><strong>总体要求</strong></p><ol><li>线上和线下二选一</li></ol></li>
<li><p><strong>线下汇报</strong></p><ol><li>汇报课前提交 PPT 和论文 PDF</li></ol></li>
<li><p><strong>线上汇报</strong></p><ol><li>提交录屏、PPT 和论文 PDF</li><li>截止日期为 <strong>12 月 15 日</strong>。</li></ol></li>
</ul>
<ul>
<li><p>网络实验实践依托平台，为课程<strong>选做内容</strong></p><ol><li>前四次任务与打榜</li></ol></li>
<li><p><strong>第一部分：前四次任务</strong></p><ul><li><strong>实验时间：</strong>9 月 14 日 00:00 至 11 月 8 日 23:59；</li><li>提交代码和实验报告</li></ul></li>
<li><p><strong>第二部分：平台打榜</strong></p><ul><li><strong>实验时间：</strong>11 月 9 日 00:00 至 2027 年 1 月 10 日 23:59；</li><li>须先完成前四次任务</li></ul></li>
</ul></div>
<div class="course-table course-table--schedule"><table><tbody>
<tr><td>第 1 讲</td><td>DCN</td><td><a href="papers/1_FatTree.pdf">参考论文</a></td><td><a href="slides/1_FatTree.pdf">FatTree</a></td></tr>
<tr><td>第 2 讲</td><td>DCN</td><td>参考论文</td><td><a href="https://evil.invalid/slides/bad.pdf">外站</a></td></tr>
</tbody></table></div>
<div class="course-table--presentations"><table><tr><td>2026年9月29日</td><td>其他同学</td><td>论文</td></tr></table></div>
</article>"""


def material(**values):
    return {
        "id": "file-1",
        "title": "ch01.pdf",
        "url": BB_ORIGIN + "/bbcswebdav/file-1",
        "description": "课件",
        "signature": "source-etag-1",
        "modified_at": "2026-09-21T14:00:00+08:00",
        "text": "监督学习的概念与模型评估。",
        "read_error": "",
        **values,
    }


def task(**values):
    return {
        "id": "lab-1",
        "title": "实验1",
        "description": "实现并比较算法",
        "url": COURSE["url"],
        "deadline": "2026-09-28T19:00:00+08:00",
        "precision": "minute",
        "submitted": None,
        **values,
    }


def snapshot(*, files=None, tasks=None, **values):
    return {
        "courses": [
            {
                "id": "ml",
                "status": "ok",
                "materials": files or [],
                "tasks": tasks or [],
                **values,
            }
        ]
    }


class CourseParserTests(unittest.IsolatedAsyncioTestCase):
    def test_deadline_requires_explicit_date_and_does_not_invent_time(self):
        self.assertEqual(
            explicit_deadline("上课日期2026-09-28；下课一周内提交"), (None, "unknown")
        )
        self.assertEqual(
            explicit_deadline("截止日期：2026年9月29日"), ("2026-09-29", "date")
        )
        self.assertEqual(
            explicit_deadline("DDL: 2026/09/29 23:59"),
            ("2026-09-29T23:59:00+08:00", "minute"),
        )
        self.assertEqual(
            explicit_deadline("截止2026-09-29 10:00 截止2026-09-30 10:00"),
            (None, "unknown"),
        )
        self.assertEqual(explicit_deadline("截止2026-99-99"), (None, "unknown"))
        self.assertIn(
            "09-29 14:00", deadline_label(task(deadline="2026-09-29 14:00:00"))
        )

    def test_blackboard_parses_folders_files_and_tasks_without_following_attempts(self):
        markup = """<ul id="courseMenuPalette_contents"><li><a href="/webapps/blackboard/content/launchLink.jsp?course_id=_37268_1&amp;tool_id=_140_1">课程通知</a></li></ul>
        <ul id="content_listContainer">
        <li id="folder"><h3><a href="/webapps/blackboard/content/listContent.jsp?course_id=_37268_1&amp;content_id=_2_1">实验</a></h3></li>
        <li id="lab1"><h3>实验1</h3><div>截止日期：2026年9月29日 23:59</div><a href="/bbcswebdav/lab.pdf">实验1.pdf</a><a href="/webapps/assessment/take/launch.jsp">开始测试</a><script>secret_random_id</script></li>
        <li id="ml1"><h3>线性模型</h3><a href="/bbcswebdav/ch03">ch03.pdf</a><a href="https://evil.invalid/a.pdf">external</a></li></ul>"""
        tasks, files, folders, notices = parse_blackboard(
            document(markup), COURSE["url"], "_37268_1"
        )
        self.assertEqual(len(tasks), 1)
        self.assertEqual(tasks[0]["deadline"], "2026-09-29T23:59:00+08:00")
        self.assertEqual(len(files), 2)
        self.assertEqual(files[0]["parent_id"], "lab1")
        self.assertNotIn("secret_random_id", tasks[0]["description"])
        self.assertEqual(len(folders), 1)
        self.assertEqual(len(notices), 1)
        self.assertFalse(any("assessment" in url for url in folders + notices))

    def test_login_page_is_not_an_empty_course(self):
        with self.assertRaisesRegex(CourseError, "session_or_access"):
            parse_blackboard(
                document("<html>统一身份认证登录</html>"), COURSE["url"], "_37268_1"
            )

    def test_algorithm_task_and_announcements(self):
        doc = document(
            '<section id="homework">每个模块讲解完成后一周提交，具体时间以 BB 平台为准。第一次作业：Ex 2.1-2.5</section><section id="courseware"><a href="ppt/a.ppt">课件</a></section>'
        )
        tasks, files = parse_algorithms(doc, DEFAULT_COURSES[0]["url"])
        self.assertEqual(len(tasks), 1)
        self.assertIsNone(tasks[0]["deadline"])
        self.assertEqual(len(files), 1)
        notices = parse_announcements(
            document(
                '<ul id="announcementList"><li id="n1"><h3>实验提交提醒</h3>截止2026-09-29 23:00</li></ul>'
            ),
            COURSE["url"],
        )
        self.assertEqual(notices[0]["id"], "notice:n1")
        self.assertEqual(notices[0]["deadline"], "2026-09-29T23:00:00+08:00")

    def test_network_tasks_use_course_year_and_label_choices_and_optional_labs(self):
        tasks, files = parse_networks(document(NETWORK_PAGE), NETWORK_URL)
        tasks = {task["id"]: task for task in tasks}
        self.assertEqual(len(tasks), 4)
        self.assertEqual(tasks["paper-online"]["deadline"], "2026-12-15")
        self.assertEqual(tasks["paper-online"]["precision"], "date")
        self.assertIsNone(tasks["paper-offline"]["deadline"])
        self.assertIn("课前提交", tasks["paper-offline"]["title"])
        self.assertTrue(
            all(
                "二选一" in tasks[key]["title"]
                for key in ("paper-online", "paper-offline")
            )
        )
        self.assertEqual(tasks["lab-tasks"]["deadline"], "2026-11-08T23:59:00+08:00")
        self.assertEqual(
            tasks["lab-leaderboard"]["deadline"], "2027-01-10T23:59:00+08:00"
        )
        self.assertTrue(
            all(
                "选做" in tasks[key]["title"]
                for key in ("lab-tasks", "lab-leaderboard")
            )
        )
        self.assertEqual(len(files), 1)
        self.assertEqual(files[0]["url"], NETWORK_URL + "slides/1_FatTree.pdf")
        self.assertEqual(explicit_deadline("截止日期为12 月 15 日"), (None, "unknown"))
        with self.assertRaisesRegex(CourseError, "course_year_missing"):
            parse_networks(document(NETWORK_PAGE.replace("（2026）", "")), NETWORK_URL)
        with self.assertRaisesRegex(CourseError, "layout_changed"):
            parse_networks(document("<article>维护中</article>"), NETWORK_URL)

    async def test_network_reader_only_downloads_slides(self):
        calls = []

        def respond(request):
            calls.append((request.method, str(request.url)))
            if str(request.url) == NETWORK_URL:
                return httpx.Response(200, text=NETWORK_PAGE)
            self.assertEqual(str(request.url), NETWORK_URL + "slides/1_FatTree.pdf")
            return httpx.Response(
                200,
                headers={"etag": "build-1", "content-type": "text/plain"},
                content=b"FatTree lecture",
            )

        client = httpx.AsyncClient(transport=httpx.MockTransport(respond))
        with (
            patch(
                "astrbot_plugin_dududa_notify.course_sources.httpx.AsyncClient",
                return_value=client,
            ),
            patch(
                "astrbot_plugin_dududa_notify.course_sources.file_text",
                return_value="FatTree lecture",
            ),
        ):
            result = await fetch_courses({"courses": [DEFAULT_COURSES[-1]]}, None)
        course = result["courses"][0]
        self.assertEqual(course["status"], "ok")
        self.assertEqual(len(course["tasks"]), 4)
        self.assertEqual(len(course["materials"]), 1)
        self.assertIsNone(course["materials"][0]["modified_at"])
        self.assertEqual(
            calls,
            [
                ("GET", NETWORK_URL),
                ("HEAD", NETWORK_URL + "slides/1_FatTree.pdf"),
                ("GET", NETWORK_URL + "slides/1_FatTree.pdf"),
            ],
        )

    async def test_network_rebuild_etag_is_not_a_file_update(self):
        calls, build, content = [], "build-1", b"lecture version one"

        def respond(request):
            calls.append(request.method)
            return httpx.Response(
                200,
                headers={
                    "etag": build,
                    "last-modified": "Sun, 27 Sep 2026 09:25:17 GMT",
                    "content-type": "text/plain",
                },
                content=content,
            )

        item = material(url=NETWORK_URL + "slides/lecture.txt")
        async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
            first = await material_snapshot(client, item, None, blackboard=False)
            cached = await material_snapshot(client, item, first, blackboard=False)
            self.assertEqual(cached, first)
            self.assertEqual(calls, ["HEAD", "GET", "HEAD"])
            build = "build-2"
            rebuilt = await material_snapshot(client, item, first, blackboard=False)
            self.assertEqual(first["signature"], rebuilt["signature"])
            self.assertNotEqual(first["http_signature"], rebuilt["http_signature"])
            self.assertIsNone(rebuilt["modified_at"])
            build, content = "build-3", b"lecture version two"
            updated = await material_snapshot(client, item, rebuilt, blackboard=False)
            self.assertNotEqual(first["signature"], updated["signature"])

    async def test_bb_redirect_never_leaks_session_to_external_host(self):
        calls = []

        def respond(request):
            calls.append(request)
            return httpx.Response(
                302, headers={"location": "https://evil.invalid/collect"}
            )

        async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
            with self.assertRaisesRegex(CourseError, "redirect_rejected"):
                await bb_get(client, COURSE["url"])
        self.assertEqual(len(calls), 1)

    async def test_unchanged_file_uses_server_metadata_not_another_download(self):
        calls = []

        def respond(request):
            calls.append(request.method)
            return httpx.Response(
                200,
                headers={
                    "etag": "v1",
                    "content-type": "text/plain",
                    "last-modified": "Mon, 21 Sep 2026 06:00:00 GMT",
                },
                content=b"normal course text",
            )

        item = material(title="slides.txt")
        async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
            first = await material_snapshot(client, item, None, blackboard=True)
            second = await material_snapshot(client, item, first, blackboard=True)
        self.assertEqual(calls, ["HEAD", "GET", "HEAD"])
        self.assertEqual(first, second)

    def test_configuration_is_explicit_and_pending_courses_stay_empty(self):
        self.assertEqual(len(validate_courses(DEFAULT_COURSES)), 7)
        self.assertEqual(
            [c["name"] for c in DEFAULT_COURSES if c["kind"] == "pending"],
            ["组合数学"],
        )
        with self.assertRaisesRegex(ValueError, "duplicate"):
            validate_courses([COURSE, COURSE])
        with self.assertRaisesRegex(ValueError, "invalid_blackboard"):
            validate_courses([{**COURSE, "url": "https://evil.invalid/"}])
        with self.assertRaisesRegex(ValueError, "unsupported_course_website"):
            validate_courses(
                [
                    {
                        **DEFAULT_COURSES[-1],
                        "url": "https://acn-2026.github.io.evil.invalid/",
                    }
                ]
            )

    def test_old_ppt_conversion_extracts_real_slide_text(self):
        from pptx import Presentation
        from pptx.util import Inches

        def convert(command, **kwargs):
            destination = Path(command[command.index("--outdir") + 1]) / "material.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_textbox(
                Inches(1), Inches(1), Inches(6), Inches(2)
            ).text = "Flooding 算法构造生成树"
            presentation.save(destination)
            return SimpleNamespace(returncode=0)

        with (
            patch(
                "astrbot_plugin_dududa_notify.course_sources.shutil.which",
                return_value="/usr/bin/libreoffice",
            ),
            patch(
                "astrbot_plugin_dududa_notify.course_sources.subprocess.run",
                side_effect=convert,
            ) as process,
        ):
            text = file_text("lecture.ppt", b"fake old-format presentation")
        self.assertIn("Flooding 算法构造生成树", text)
        self.assertIn("--headless", process.call_args.args[0])
        self.assertEqual(process.call_args.kwargs["timeout"], 60)


class CourseWorkflowTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.now = datetime(2026, 9, 27, 19, 0, tzinfo=CHINA_TIME)
        clock = patch(
            "astrbot_plugin_dududa_notify.course_tasks.current_time",
            side_effect=lambda: self.now,
        )
        clock.start()
        self.addCleanup(clock.stop)
        self.query = AsyncMock(
            return_value=snapshot(files=[material()], tasks=[task()])
        )
        self.generate = AsyncMock(return_value=("- 监督学习\n- 模型评估", "test-model"))
        self.render = AsyncMock(return_value=b"png-test")
        self.text, self.image = (
            AsyncMock(return_value="text-ack"),
            AsyncMock(return_value="image-ack"),
        )
        self.service = self.make_service()
        self.addCleanup(lambda: self.service.close())

    def make_service(self):
        return CourseService(
            self.directory.name,
            RECIPIENT,
            [COURSE, DEFAULT_COURSES[-2]],
            self.query,
            self.generate,
            self.render,
            self.text,
            self.image,
            summary_hour=20,
        )

    async def test_initial_old_slides_are_not_todays_uploads(self):
        await self.service.poll()
        row = self.service.db.execute("SELECT * FROM materials").fetchone()
        self.assertIsNone(row["changed_on"])
        await self.service.daily()
        self.generate.assert_not_awaited()
        report = self.render.await_args.args[0]
        self.assertIn("课件：1 份；今日新增/更新 0 份", report["sections"][0]["lines"])
        self.assertIn("待补充", report["sections"][1]["lines"][0])
        self.image.assert_awaited_once()
        await self.service.daily()
        self.image.assert_awaited_once()

    def test_cross_year_deadline_label_includes_the_year(self):
        self.assertEqual(
            deadline_label(task(deadline="2027-01-10T23:59:00+08:00")),
            "2027-01-10 23:59",
        )

    async def test_deployment_metadata_only_change_does_not_queue_another_summary(self):
        original = material(modified_at=None, http_signature="build-1")
        self.service.sync(snapshot(files=[original]))
        self.service.sync(snapshot(files=[{**original, "http_signature": "build-2"}]))
        row = self.service.db.execute("SELECT * FROM materials").fetchone()
        self.assertEqual(row["revision"], 1)
        self.assertIsNone(row["changed_on"])
        self.service.sync(
            snapshot(
                files=[
                    {
                        **original,
                        "http_signature": "build-3",
                        "signature": "changed-pdf-bytes",
                    }
                ]
            )
        )
        row = self.service.db.execute("SELECT * FROM materials").fetchone()
        self.assertEqual(row["revision"], 2)
        self.assertEqual(row["changed_on"], self.now.date().isoformat())

    async def test_daily_at_20_and_new_files_or_revisions_are_summarized_once(self):
        await self.service.poll()
        await self.service.daily()
        self.query.return_value = snapshot(
            files=[material(), material(id="file-2", title="决策树.pdf")]
        )
        await self.service.poll()
        await self.service.daily()
        self.generate.assert_not_awaited()
        self.now += timedelta(hours=1)
        await self.service.daily()
        self.generate.assert_awaited_once()
        self.assertEqual(self.image.await_count, 2)
        await self.service.daily()
        self.assertEqual(self.image.await_count, 2)
        self.query.return_value = snapshot(
            files=[
                material(),
                material(
                    id="file-2", title="决策树.pdf", signature="v2", text="新内容"
                ),
            ]
        )
        await self.service.poll()
        await self.service.daily()
        self.assertEqual(self.generate.await_count, 2)
        self.assertEqual(self.image.await_count, 3)

    async def test_empty_course_followed_by_first_file_counts_as_new(self):
        self.service.sync(snapshot())
        self.service.sync(snapshot(files=[material()]))
        self.assertEqual(
            self.service.db.execute("SELECT changed_on FROM materials").fetchone()[0],
            "2026-09-27",
        )

    async def test_first_connection_after_20_sends_one_report(self):
        self.now += timedelta(hours=1)
        await self.service.poll()
        await self.service.daily()
        await self.service.daily()
        self.image.assert_awaited_once()
        self.assertEqual(
            self.service.db.execute("SELECT kind FROM reports").fetchone()[0], "daily"
        )

    async def test_initial_file_modified_today_is_in_daily_summary(self):
        self.query.return_value = snapshot(
            files=[material(modified_at=self.now.isoformat())]
        )
        await self.service.poll()
        self.now += timedelta(hours=1)
        await self.service.daily()
        self.generate.assert_awaited_once()

    async def test_reminders_only_once_per_threshold_and_after_restart(self):
        await self.service.poll()
        first = self.text.await_count
        await self.service.poll()
        self.assertEqual(self.text.await_count, first)
        self.now += timedelta(hours=18)
        await self.service.poll()
        self.assertIn("6 小时", self.text.await_args.args[0])
        self.service.close()
        self.service = self.make_service()
        count = self.text.await_count
        await self.service.poll()
        self.assertEqual(self.text.await_count, count)
        self.now += timedelta(hours=5)
        await self.service.poll()
        self.assertIn("1 小时", self.text.await_args.args[0])

    async def test_done_or_school_submitted_task_has_no_reminders(self):
        await self.service.poll()
        await self.service.command("完成 C1")
        self.now += timedelta(hours=23)
        count = self.text.await_count
        await self.service.poll()
        self.assertEqual(self.text.await_count, count)
        await self.service.command("恢复 C1")
        self.query.return_value = snapshot(tasks=[task(submitted=True)])
        await self.service.poll()
        self.assertEqual(self.text.await_count, count)

    async def test_date_only_and_unknown_deadlines_are_not_fake_midnight(self):
        self.assertEqual(
            self.service.reminder_stage(task(deadline="2026-09-27", precision="date")),
            "date",
        )
        self.assertIsNone(self.service.reminder_stage(task(deadline=None)))
        self.assertIn(
            "具体时刻未注明",
            deadline_label(task(deadline="2026-09-27", precision="date")),
        )

    async def test_failed_source_keeps_cache_and_alerts_once(self):
        await self.service.poll()
        self.query.return_value = snapshot(status="error", error="ReadTimeout")
        for _ in range(4):
            await self.service.poll()
        errors = [c for c in self.text.await_args_list if "课程同步异常" in c.args[0]]
        self.assertEqual(len(errors), 1)
        self.assertIn("缓存", await self.service.command("状态"))
        self.assertEqual(
            self.service.db.execute("SELECT active FROM materials").fetchone()[0], 1
        )
        self.query.return_value = snapshot(files=[material()], tasks=[task()])
        await self.service.poll()
        self.assertIn("恢复", self.text.await_args.args[0])

    async def test_image_failure_retries_image_without_resending_text_or_regenerating(
        self,
    ):
        await self.service.poll()
        self.image.side_effect = RuntimeError("QQ offline")
        with self.assertRaises(RuntimeError):
            await self.service.daily()
        count = self.text.await_count
        self.image.side_effect = None
        await self.service.daily()
        self.assertEqual(self.text.await_count, count)
        self.render.assert_awaited_once()

    async def test_manual_deadline_is_local_and_not_overwritten_by_poll(self):
        await self.service.poll()
        result = await self.service.command("截止 C1 2026-09-30 23:59")
        self.assertIn("不修改学校", result)
        await self.service.poll()
        row = self.service.db.execute("SELECT * FROM tasks").fetchone()
        self.assertEqual(
            self.service.task_payload(row)["deadline"], "2026-09-30T23:59:00+08:00"
        )
        self.assertTrue(
            all(
                call.args[0]["operation"] == "courses"
                for call in self.query.await_args_list
            )
        )

    async def test_daily_with_failed_source_gets_one_correction_on_recovery(self):
        self.now += timedelta(hours=1)
        self.query.return_value = snapshot(status="error", error="database unavailable")
        await self.service.poll()
        await self.service.daily()
        self.assertEqual(
            self.render.await_args.args[0]["sections"][0]["status"], "error"
        )
        self.query.return_value = snapshot(tasks=[task()], files=[material()])
        await self.service.poll()
        await self.service.daily()
        self.assertEqual(self.image.await_count, 2)
        self.assertIn("修正版", self.render.await_args.args[0]["title"])
        self.assertEqual(self.render.await_args.args[0]["sections"][0]["status"], "ok")
        self.service.close()
        self.service = self.make_service()
        await self.service.daily()
        self.assertEqual(self.image.await_count, 2)

    async def test_today_legacy_report_without_status_metadata_is_corrected(self):
        self.now += timedelta(hours=1)
        self.query.return_value = snapshot(status="error", error="database unavailable")
        await self.service.poll()
        await self.service.daily()
        row = self.service.db.execute("SELECT * FROM reports").fetchone()
        old = json.loads(row["payload"])
        for section in old["sections"]:
            section.pop("status")
            section.pop("course_id")
        with self.service.db:
            self.service.db.execute(
                "UPDATE reports SET payload=? WHERE id=?", (json.dumps(old), row["id"])
            )
        self.query.return_value = snapshot()
        await self.service.poll()
        await self.service.daily()
        self.assertEqual(self.image.await_count, 2)
        self.assertIn("修正版", self.render.await_args.args[0]["title"])

    async def test_correction_image_retry_does_not_create_another_report(self):
        self.now += timedelta(hours=1)
        self.query.return_value = snapshot(status="error", error="database unavailable")
        await self.service.poll()
        await self.service.daily()
        self.query.return_value = snapshot()
        await self.service.poll()
        self.image.side_effect = RuntimeError("QQ failure")
        with self.assertRaises(RuntimeError):
            await self.service.daily()
        texts = self.text.await_count
        self.image.side_effect = None
        await self.service.daily()
        self.assertEqual(self.text.await_count, texts)
        self.assertEqual(
            self.service.db.execute("SELECT COUNT(*) FROM reports").fetchone()[0], 2
        )

    async def test_graduate_report_distinguishes_school_submission_from_local_done(
        self,
    ):
        self.service.courses = [DEFAULT_COURSES[3]]
        self.service.sync(
            {
                "courses": [
                    {
                        "id": "ai",
                        "status": "ok",
                        "tasks": [
                            task(id="submitted", submitted=True),
                            task(id="pending", submitted=False),
                        ],
                        "materials": [],
                    }
                ]
            }
        )
        with self.service.db:
            self.service.db.execute("UPDATE tasks SET done=1 WHERE remote_id='pending'")
        report, _ = await self.service.build_report(summarize=False)
        self.assertIn(
            "研究生系统：待提交 1 份；已核实已提交 1 份", report["sections"][0]["lines"]
        )

    async def test_large_material_all_chunks_are_used_and_unreadable_material_is_not_fabricated(
        self,
    ):
        self.query.return_value = snapshot(
            files=[material(text="a" * 36010, modified_at=self.now.isoformat())]
        )
        await self.service.poll()
        row = self.service.db.execute("SELECT * FROM materials").fetchone()
        await self.service.summarize(row)
        self.assertEqual(self.generate.await_count, 4)
        self.query.return_value = snapshot(
            files=[
                material(text="", read_error="扫描件", modified_at=self.now.isoformat())
            ]
        )
        await self.service.poll()
        row = self.service.db.execute("SELECT * FROM materials").fetchone()
        self.assertEqual(await self.service.summarize(row), "扫描件")
        self.assertEqual(self.generate.await_count, 4)


class CourseImageAndAccessTests(unittest.IsolatedAsyncioTestCase):
    async def test_report_png_nonblank_and_contains_all_list_rows(self):
        from PIL import Image

        report = {
            "title": "全课程任务与课件日报",
            "subtitle": "2026-09-27 · 每日20:00",
            "sections": [
                {
                    "title": c["name"],
                    "lines": [
                        "实验任务：" + "很长的中文标题" * 30,
                        "DDL：2026-09-29 23:59",
                    ],
                }
                for c in DEFAULT_COURSES
            ],
            "notes": ["按当天新增/更新汇总"],
        }
        png = render_report(report)
        image = Image.open(io.BytesIO(png))
        self.assertEqual(image.width, 1120)
        self.assertGreater(image.height, 2500)
        self.assertGreater(len(image.getcolors(image.width * image.height)), 20)
        self.assertIn("7. 高级计算机网络", report_text(report))

    async def test_only_private_owner_can_update_completion(self):
        plugin = load_plugin()(SimpleNamespace(), {})
        plugin.platform_id = "qq"
        plugin.courses = SimpleNamespace(
            owner=RECIPIENT, command=AsyncMock(return_value="")
        )
        event = SimpleNamespace(
            is_private_chat=lambda: True,
            get_sender_id=lambda: RECIPIENT,
            get_self_id=lambda: "bot",
            get_platform_id=lambda: "qq",
            message_str="课程 完成 C1",
            stop_event=Mock(),
        )
        for changes in (
            {"get_sender_id": lambda: "other"},
            {"is_private_chat": lambda: False},
            {"get_platform_id": lambda: "other"},
            {"get_self_id": lambda: RECIPIENT},
            {"message_str": "引用：课程 完成 C1"},
        ):
            await plugin.course_message(SimpleNamespace(**{**vars(event), **changes}))
        plugin.courses.command.assert_not_awaited()
        await plugin.course_message(event)
        plugin.courses.command.assert_awaited_once_with("完成 C1")

    async def test_private_image_needs_receipt_and_enforces_owner(self):
        action = AsyncMock(return_value={"message_id": "img-1"})
        platform = SimpleNamespace(
            meta=lambda: SimpleNamespace(id="qq", name="aiocqhttp"),
            get_client=lambda: SimpleNamespace(call_action=action),
        )
        plugin = load_plugin()(
            SimpleNamespace(
                platform_manager=SimpleNamespace(get_insts=lambda: [platform])
            ),
            {},
        )
        plugin.platform_id, plugin.recipients = "qq", (RECIPIENT,)
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "report.png"
            path.write_bytes(b"png")
            await plugin._send_private_image(RECIPIENT, path)
            self.assertEqual(action.await_args.kwargs["message"][0]["type"], "image")
            with self.assertRaisesRegex(ValueError, "not_authorized"):
                await plugin._send_private_image("other", path)
            action.return_value = {}
            with self.assertRaisesRegex(RuntimeError, "not_acknowledged"):
                await plugin._send_private_image(RECIPIENT, path)


if __name__ == "__main__":
    unittest.main()
